from datetime import datetime
import os
import struct
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side

from merge_workbooks import (
    copy_cell,
    find_workbooks,
    get_header_rows,
    get_input_folder,
    load_workbook_file,
    merge_workbooks,
    natural_sort_key,
    save_workbook_atomic,
)


class MergeWorkbooksTests(unittest.TestCase):
    def test_natural_sort_orders_numeric_names(self) -> None:
        names = [
            "13-16.xlsx",
            "0-9.xlsx",
            "20-24.xlsx",
            "9-13.xlsx",
            "16-20.xlsx",
        ]

        self.assertEqual(
            sorted(names, key=natural_sort_key),
            [
                "0-9.xlsx",
                "9-13.xlsx",
                "13-16.xlsx",
                "16-20.xlsx",
                "20-24.xlsx",
            ],
        )

    def test_find_workbooks_recursively(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            nested_folder = folder / "nested"
            nested_folder.mkdir()
            output_file = folder / "合并结果.xlsx"

            included_files = [
                folder / "root.XLSX",
                nested_folder / "child.xlsm",
                nested_folder / "legacy.XLS",
                nested_folder / "table.CSV",
            ]
            excluded_files = [
                nested_folder / "~$temporary.xlsx",
                nested_folder / "~$temporary.xls",
                nested_folder / "~$temporary.csv",
                folder / ".合并结果-previous.xlsx",
                folder / ".合并结果-current.tmp",
                folder / "notes.txt",
                output_file,
            ]
            for path in included_files + excluded_files:
                path.touch()

            self.assertEqual(
                find_workbooks(folder, output_file),
                [
                    nested_folder / "child.xlsm",
                    nested_folder / "legacy.XLS",
                    nested_folder / "table.CSV",
                    folder / "root.XLSX",
                ],
            )

    def test_csv_is_loaded_as_a_single_worksheet(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            csv_file = Path(temporary_directory) / "sales.csv"
            csv_file.write_text("Name,Amount\nAlice,12\nBob,8\n", encoding="utf-8")

            workbook = load_workbook_file(csv_file)
            try:
                self.assertEqual(workbook.sheetnames, ["sales"])
                self.assertEqual(
                    list(workbook["sales"].values),
                    [("Name", "Amount"), ("Alice", "12"), ("Bob", "8")],
                )
            finally:
                workbook.close()

    def test_csv_handles_encodings_delimiters_and_quoted_newlines(self) -> None:
        cases = [
            ("utf-8-sig", ",", "张三", "中文\n换行"),
            ("gb18030", ";", "李四", "报价;备注"),
            ("utf-8", "\t", "Alice", "quoted\ttext"),
            ("utf-8", "|", "Bob", "quoted|text"),
        ]
        for encoding, delimiter, name, note in cases:
            with self.subTest(encoding=encoding, delimiter=delimiter):
                with tempfile.TemporaryDirectory() as temporary_directory:
                    source_file = Path(temporary_directory) / "表格.csv"
                    source_file.write_text(
                        f'姓名{delimiter}备注\r\n{name}{delimiter}"{note}"\r\n',
                        encoding=encoding,
                        newline="",
                    )
                    workbook = load_workbook_file(source_file)
                    try:
                        self.assertEqual(workbook.sheetnames, ["表格"])
                        self.assertEqual(
                            list(workbook.active.values),
                            [("姓名", "备注"), (name, note)],
                        )
                    finally:
                        workbook.close()

    def test_csv_encoding_retry_does_not_duplicate_rows(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            source_file = Path(temporary_directory) / "Data.csv"
            source_file.write_text(
                "Name,Value\n" + "ASCII,1\n" * 3000 + "中文,2\n",
                encoding="gb18030",
            )
            workbook = load_workbook_file(source_file)
            try:
                self.assertEqual(workbook.active.max_row, 3002)
                self.assertEqual(workbook.active["A3002"].value, "中文")
            finally:
                workbook.close()

    def test_csv_merge_preserves_text_and_skips_blank_rows(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            source_file = folder / "Data.csv"
            source_file.write_text(
                "Code,Note\n0012,=A2\n,\n\n0034,#DIV/0!\n",
                encoding="utf-8",
            )
            output_file = folder / "output.xlsx"
            merge_workbooks([source_file], output_file)
            workbook = load_workbook(output_file)
            try:
                self.assertEqual(
                    list(workbook["Data"].values),
                    [("Code", "Note"), ("0012", "=A2"), ("0034", "#DIV/0!")],
                )
                self.assertEqual(workbook["Data"]["B2"].data_type, "s")
                self.assertEqual(workbook["Data"]["B3"].data_type, "s")
            finally:
                workbook.close()

    def test_empty_csv_does_not_create_a_source_sheet_in_output(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            source_file = folder / "empty.csv"
            source_file.write_text(",,\n\n,,\n", encoding="utf-8")
            output_file = folder / "output.xlsx"
            self.assertEqual(merge_workbooks([source_file], output_file), (1, 0))
            workbook = load_workbook(output_file)
            try:
                self.assertEqual(workbook.sheetnames, ["合并结果"])
            finally:
                workbook.close()

    def test_csv_uses_a_valid_excel_sheet_title(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            source_file = Path(temporary_directory) / ("a" * 32 + "[data].csv")
            source_file.write_text("header\nvalue\n", encoding="utf-8")
            workbook = load_workbook_file(source_file)
            try:
                self.assertEqual(workbook.sheetnames, ["a" * 31])
            finally:
                workbook.close()

    def test_xls_preserves_values_number_formats_and_merged_cells(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            source_file = folder / "legacy.xls"
            source_file.write_bytes(self.legacy_workbook_bytes())
            output_file = folder / "output.xlsx"
            self.assertEqual(merge_workbooks([source_file], output_file), (1, 2))
            workbook = load_workbook(output_file)
            try:
                self.assertEqual(workbook.sheetnames, ["Data", "Other"])
                worksheet = workbook["Data"]
                self.assertEqual(worksheet["A2"].value, "legacy")
                self.assertEqual(worksheet["B2"].value, 12.5)
                self.assertEqual(worksheet["C2"].value, datetime(2026, 9, 28))
                self.assertEqual(worksheet["C2"].number_format, "m/d/yy")
                self.assertIs(worksheet["D2"].value, True)
                self.assertEqual(worksheet["E2"].value, "#DIV/0!")
                self.assertEqual(worksheet["E2"].data_type, "e")
                self.assertEqual(worksheet["F2"].value, "=literal")
                self.assertEqual(worksheet["F2"].data_type, "s")
                self.assertEqual(worksheet["G2"].value, 42)
                self.assertEqual(str(worksheet.merged_cells), "A3:B4")
                self.assertEqual(worksheet.column_dimensions["B"].width, 20)
                self.assertEqual(worksheet.row_dimensions[2].height, 30)
                self.assertEqual(workbook["Other"]["A1"].value, "separate")
            finally:
                workbook.close()

    def test_missing_xlrd_only_blocks_xls_with_install_hint(self) -> None:
        with patch("merge_workbooks.xlrd", None):
            with self.assertRaisesRegex(RuntimeError, "xlrd.*requirements.txt"):
                load_workbook_file(Path("legacy.xls"))

    def test_read_errors_include_source_path(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            for suffix in (".xls", ".xlsx", ".xlsm", ".csv"):
                with self.subTest(suffix=suffix):
                    source_file = Path(temporary_directory) / f"broken{suffix}"
                    source_file.write_bytes(b"\xff\xff\xff")
                    with self.assertRaises(RuntimeError) as captured:
                        load_workbook_file(source_file)
                    self.assertIn(str(source_file), str(captured.exception))

    def test_mixed_formats_merge_same_named_sheets_in_source_order(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            legacy_file = folder / "01.xls"
            legacy_file.write_bytes(self.legacy_workbook_bytes())
            for suffix in ("xlsx", "xlsm"):
                workbook = Workbook()
                worksheet = workbook.active
                worksheet.title = "Data"
                worksheet.append(["Name", "Value"])
                worksheet.append([suffix, 20])
                workbook.save(folder / f"02.{suffix}")
                workbook.close()
            csv_file = folder / "Data.csv"
            csv_file.write_text("Name,Value\ncsv,30\n", encoding="utf-8")
            output_file = folder / "output.xlsx"
            for workers in (1, 2):
                with self.subTest(workers=workers):
                    files = find_workbooks(folder, output_file)
                    self.assertEqual(
                        merge_workbooks(files, output_file, workers=workers),
                        (4, 5),
                    )
                    workbook = load_workbook(output_file)
                    try:
                        self.assertEqual(
                            [cell.value for cell in workbook["Data"]["A"]],
                            ["Name", "legacy", "merged", None, "xlsm", "xlsx", "csv"],
                        )
                        self.assertEqual(workbook.sheetnames, ["Data", "Other"])
                    finally:
                        workbook.close()

    @staticmethod
    def legacy_workbook_bytes(
        number_format_id: int | None = None,
        style_index: int = 1,
    ) -> bytes:
        fixture_file = Path(__file__).parent / "fixtures" / "legacy_workbook.hex"
        contents = bytearray(bytes.fromhex(fixture_file.read_text(encoding="ascii")))
        if number_format_id is not None:
            position = 0
            current_style = 0
            while position + 4 <= len(contents):
                record_code, length = struct.unpack_from("<HH", contents, position)
                if record_code == 0x00E0:
                    if current_style == style_index:
                        struct.pack_into("<H", contents, position + 6, number_format_id)
                        break
                    current_style += 1
                position += 4 + length
            else:
                raise ValueError("Fixture does not contain the requested style")
        return bytes(contents)

    def test_xls_localized_dates_with_unknown_format_strings_remain_dates(self) -> None:
        for number_format_id in (27, 31, 50, 71):
            with self.subTest(number_format_id=number_format_id):
                with tempfile.TemporaryDirectory() as temporary_directory:
                    folder = Path(temporary_directory)
                    source_file = folder / "localized.xls"
                    source_file.write_bytes(self.legacy_workbook_bytes(number_format_id))
                    output_file = folder / "output.xlsx"
                    merge_workbooks([source_file], output_file)
                    workbook = load_workbook(output_file)
                    try:
                        self.assertEqual(workbook["Data"]["C2"].value, datetime(2026, 9, 28))
                        self.assertEqual(
                            workbook["Data"]["C2"].number_format,
                            "yyyy-mm-dd hh:mm:ss",
                        )
                        self.assertTrue(
                            all(isinstance(value, str) for value in workbook._number_formats),
                        )
                    finally:
                        workbook.close()

    def test_xls_localized_numbers_with_unknown_format_strings_use_general(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            source_file = folder / "localized.xls"
            source_file.write_bytes(self.legacy_workbook_bytes(59, style_index=0))
            output_file = folder / "output.xlsx"
            merge_workbooks([source_file], output_file)
            workbook = load_workbook(output_file)
            try:
                self.assertEqual(workbook["Data"]["B2"].value, 12.5)
                self.assertEqual(workbook["Data"]["B2"].number_format, "General")
                self.assertEqual(workbook["Data"]["C2"].value, datetime(2026, 9, 28))
            finally:
                workbook.close()

    def test_xls_missing_number_format_definition_does_not_block_merging(self) -> None:
        import xlrd

        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            source_file = folder / "legacy.xls"
            source_book = xlrd.open_workbook(
                file_contents=self.legacy_workbook_bytes(),
                formatting_info=True,
            )
            source_book.format_map.pop(0)
            source_book.format_map.pop(14)
            output_file = folder / "output.xlsx"
            with patch("merge_workbooks.xlrd.open_workbook", return_value=source_book):
                merge_workbooks([source_file], output_file)
            workbook = load_workbook(output_file)
            try:
                self.assertEqual(workbook["Data"]["B2"].value, 12.5)
                self.assertEqual(workbook["Data"]["C2"].value, datetime(2026, 9, 28))
            finally:
                workbook.close()

    def test_copy_cell_uses_typed_defaults_for_invalid_number_formats(self) -> None:
        for number_format in (None, "", "   ", 42):
            with self.subTest(number_format=number_format):
                with tempfile.TemporaryDirectory() as temporary_directory:
                    source_workbook = Workbook()
                    output_workbook = Workbook()
                    try:
                        source_workbook.active["A1"] = 12.5
                        source_workbook.active["A2"] = datetime(2026, 9, 28)
                        style_cache: dict[int, object] = {}
                        for row in (1, 2):
                            source_cell = source_workbook.active.cell(row=row, column=1)
                            source_cell.number_format = number_format
                            copy_cell(
                                source_cell,
                                output_workbook.active.cell(row=row, column=1),
                                style_cache,
                            )
                        output_file = Path(temporary_directory) / "output.xlsx"
                        save_workbook_atomic(output_workbook, output_file)
                        result_workbook = load_workbook(output_file)
                        try:
                            self.assertEqual(result_workbook.active["A1"].value, 12.5)
                            self.assertEqual(result_workbook.active["A1"].number_format, "General")
                            self.assertEqual(
                                result_workbook.active["A2"].value,
                                datetime(2026, 9, 28),
                            )
                            self.assertEqual(
                                result_workbook.active["A2"].number_format,
                                "yyyy-mm-dd hh:mm:ss",
                            )
                        finally:
                            result_workbook.close()
                    finally:
                        source_workbook.close()
                        output_workbook.close()

    def test_get_input_folder_prompts_when_argument_is_missing(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            with patch("builtins.input", return_value=f'"{temporary_directory}"'):
                self.assertEqual(
                    get_input_folder(None),
                    Path(temporary_directory).resolve(),
                )

    def test_get_header_rows_prompts_and_retries_invalid_values(self) -> None:
        with patch("builtins.input", side_effect=["invalid", "-1", "2"]):
            with patch("builtins.print"):
                self.assertEqual(get_header_rows(None, prompt_when_missing=True), 2)

        self.assertEqual(get_header_rows(None, prompt_when_missing=False), 1)

    def test_merge_skips_selected_number_of_header_rows(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            source_files = [folder / "first.xlsx", folder / "second.xlsx"]
            output_file = folder / "output.xlsx"

            for source_file, label in zip(source_files, ["first", "second"]):
                workbook = Workbook()
                worksheet = workbook.active
                worksheet.title = "Data"
                worksheet.append([f"{label} title 1"])
                worksheet.append([f"{label} title 2"])
                worksheet.append([f"{label} data"])
                workbook.save(source_file)
                workbook.close()

            progress_updates: list[float] = []
            merge_workbooks(
                source_files,
                output_file,
                header_rows=2,
                workers=2,
                progress=lambda fraction, _label: progress_updates.append(fraction),
            )

            merged_workbook = load_workbook(output_file, read_only=True)
            try:
                values = [row[0] for row in merged_workbook["Data"].values]
                self.assertEqual(
                    values,
                    ["first title 1", "first title 2", "first data", "second data"],
                )
                self.assertEqual(progress_updates[-1], 1.0)
                self.assertEqual(progress_updates, sorted(progress_updates))
            finally:
                merged_workbook.close()

    def test_merge_removes_empty_rows_and_translates_formulas(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            first_file = folder / "first.xlsx"
            second_file = folder / "second.xlsx"
            output_file = folder / "output.xlsx"

            for source_file, first_value, second_value in [
                (first_file, 10, 20),
                (second_file, 50, 60),
            ]:
                workbook = Workbook()
                worksheet = workbook.active
                worksheet.title = "Data"
                worksheet.append(["Name", "Value", "Total"])
                worksheet.append(["first", first_value, None])
                worksheet["C2"] = "=A2+B2"
                worksheet.cell(row=3, column=1).value = None
                worksheet.append(["second", second_value, None])
                worksheet["C4"] = "=A4+B4"
                workbook.save(source_file)
                workbook.close()

            merge_workbooks(
                [first_file, second_file],
                output_file,
                header_rows=1,
            )

            merged_workbook = load_workbook(output_file, data_only=False)
            try:
                worksheet = merged_workbook["Data"]
                self.assertEqual(worksheet.max_row, 5)
                self.assertEqual(
                    [worksheet.cell(row=row, column=1).value for row in range(1, 6)],
                    ["Name", "first", "second", "first", "second"],
                )
                self.assertEqual(
                    [worksheet.cell(row=row, column=3).value for row in range(2, 6)],
                    ["=A2+B2", "=A3+B3", "=A4+B4", "=A5+B5"],
                )
            finally:
                merged_workbook.close()

    def test_merge_removes_styled_empty_rows(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            source_file = Path(temporary_directory) / "source.xlsx"
            output_file = Path(temporary_directory) / "output.xlsx"

            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "Data"
            worksheet["A1"] = "kept"
            worksheet["A2"].fill = PatternFill(fill_type="solid", fgColor="FFFFFF")
            worksheet["A2"].value = None
            worksheet["A3"] = "also kept"
            workbook.save(source_file)
            workbook.close()

            merge_workbooks([source_file], output_file, header_rows=0)

            merged_workbook = load_workbook(output_file, data_only=False)
            try:
                worksheet = merged_workbook["Data"]
                self.assertEqual(worksheet.max_row, 2)
                self.assertEqual(
                    [cell.value for cell in worksheet["A"]],
                    ["kept", "also kept"],
                )
            finally:
                merged_workbook.close()

    def test_merge_translates_formulas_from_rows_after_empty_prefix(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            source_file = folder / "source.xlsx"
            output_file = folder / "output.xlsx"

            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "Data"
            worksheet["A1"] = "header 1"
            worksheet["A2"] = "header 2"
            for row, serial in [(11, 9), (71, 69)]:
                worksheet.cell(row=row, column=1).value = serial
                worksheet.cell(row=row, column=2).value = row
                worksheet.cell(row=row, column=3).value = f"=A{row}-B{row}"
            worksheet["D11"] = "=A1-B2"
            worksheet["D71"] = "=A1-B2"
            workbook.save(source_file)
            workbook.close()

            merge_workbooks([source_file], output_file, header_rows=0)

            merged_workbook = load_workbook(output_file, data_only=False)
            try:
                worksheet = merged_workbook["Data"]
                self.assertEqual(worksheet["A3"].value, 9)
                self.assertEqual(worksheet["C3"].value, "=A3-B3")
                self.assertEqual(worksheet["A4"].value, 69)
                self.assertEqual(worksheet["C4"].value, "=A4-B4")
                self.assertEqual(worksheet["D3"].value, "=A1-B2")
                self.assertEqual(worksheet["D4"].value, "=A1-B2")
                self.assertNotIn("#REF!", worksheet["C3"].value)
                self.assertNotIn("#REF!", worksheet["C4"].value)
            finally:
                merged_workbook.close()

    def test_atomic_save_preserves_existing_output_after_failure(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_file = Path(temporary_directory) / "output.xlsx"
            output_file.write_bytes(b"existing output")
            workbook = Workbook()

            try:
                with patch(
                    "openpyxl.writer.excel.ExcelWriter.write_data",
                    side_effect=RuntimeError("save failed"),
                ):
                    with self.assertRaises(RuntimeError):
                        save_workbook_atomic(workbook, output_file)
                self.assertEqual(output_file.read_bytes(), b"existing output")
                self.assertEqual(list(output_file.parent.glob(".output-*")), [])
            finally:
                workbook.close()

    def test_failed_serialization_closes_archive_and_preserves_original_error(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_file = Path(temporary_directory) / "output.xlsx"
            output_file.write_bytes(b"existing output")
            workbook = Workbook()
            original_error = ValueError("worksheet serialization failed")
            try:
                with patch(
                    "openpyxl.writer.excel.ExcelWriter.write_data",
                    side_effect=original_error,
                ):
                    with self.assertRaises(ValueError) as captured:
                        save_workbook_atomic(workbook, output_file)
                self.assertIs(captured.exception, original_error)
                self.assertEqual(output_file.read_bytes(), b"existing output")
                self.assertEqual(list(output_file.parent.glob(".output-*")), [])
            finally:
                workbook.close()

    def test_cleanup_failure_does_not_mask_serialization_error(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_file = Path(temporary_directory) / "output.xlsx"
            workbook = Workbook()
            original_error = ValueError("worksheet serialization failed")
            try:
                with patch(
                    "openpyxl.writer.excel.ExcelWriter.write_data",
                    side_effect=original_error,
                ), patch.object(Path, "unlink", side_effect=PermissionError("locked")):
                    with self.assertRaises(ValueError) as captured:
                        save_workbook_atomic(workbook, output_file)
                self.assertIs(captured.exception, original_error)
            finally:
                workbook.close()

    def test_atomic_save_retries_a_transient_windows_sharing_violation(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_file = Path(temporary_directory) / "output.xlsx"
            output_file.write_bytes(b"existing output")
            workbook = Workbook()
            workbook.active["A1"] = "new output"
            sharing_error = PermissionError("temporarily locked")
            sharing_error.winerror = 32
            original_replace = Path.replace
            attempts: list[Path] = []

            def replace_after_unlock(source: Path, target: Path) -> Path:
                attempts.append(source)
                if len(attempts) == 1:
                    raise sharing_error
                return original_replace(source, target)

            try:
                with patch.object(
                    Path,
                    "replace",
                    autospec=True,
                    side_effect=replace_after_unlock,
                ), patch("merge_workbooks.time.sleep") as sleep:
                    save_workbook_atomic(workbook, output_file)
                self.assertEqual(len(attempts), 2)
                sleep.assert_called_once_with(0.1)
                merged_workbook = load_workbook(output_file)
                try:
                    self.assertEqual(merged_workbook.active["A1"].value, "new output")
                finally:
                    merged_workbook.close()
                self.assertEqual(list(output_file.parent.glob(".output-*")), [])
            finally:
                workbook.close()

    def test_atomic_save_preserves_completed_result_when_output_stays_locked(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_file = Path(temporary_directory) / "output.xlsx"
            output_file.write_bytes(b"existing output")
            workbook = Workbook()
            workbook.active["A1"] = "new output"
            sharing_error = PermissionError("output locked")
            sharing_error.winerror = 32
            try:
                with patch.object(
                    Path,
                    "replace",
                    side_effect=sharing_error,
                ) as replace, patch("merge_workbooks.time.sleep") as sleep:
                    with self.assertRaises(RuntimeError) as captured:
                        save_workbook_atomic(workbook, output_file)
                self.assertEqual(replace.call_count, 6)
                self.assertEqual(sleep.call_count, 5)
                self.assertIs(captured.exception.__cause__, sharing_error)
                self.assertEqual(output_file.read_bytes(), b"existing output")
                temporary_files = list(output_file.parent.glob(".output-*.tmp"))
                self.assertEqual(len(temporary_files), 1)
                self.assertIn(str(temporary_files[0]), str(captured.exception))
                self.assertIn(str(output_file), str(captured.exception))
                with temporary_files[0].open("rb") as saved_file:
                    saved_workbook = load_workbook(saved_file)
                    try:
                        self.assertEqual(saved_workbook.active["A1"].value, "new output")
                    finally:
                        saved_workbook.close()
                self.assertEqual(find_workbooks(output_file.parent, output_file), [])
            finally:
                workbook.close()

    def test_atomic_save_does_not_retry_unrelated_permission_errors(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_file = Path(temporary_directory) / "output.xlsx"
            workbook = Workbook()
            try:
                with patch.object(
                    Path,
                    "replace",
                    side_effect=PermissionError("access denied"),
                ) as replace, patch("merge_workbooks.time.sleep") as sleep:
                    with self.assertRaises(RuntimeError):
                        save_workbook_atomic(workbook, output_file)
                replace.assert_called_once()
                sleep.assert_not_called()
            finally:
                workbook.close()

    @unittest.skipUnless(os.name == "nt", "Requires Windows file sharing semantics")
    def test_atomic_save_handles_an_actual_windows_file_lock(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_file = Path(temporary_directory) / "output.xlsx"
            output_file.write_bytes(b"existing output")
            workbook = Workbook()
            workbook.active["A1"] = "new output"
            try:
                with output_file.open("rb") as locked_file:
                    with patch("merge_workbooks.time.sleep"):
                        with self.assertRaises(RuntimeError) as captured:
                            save_workbook_atomic(workbook, output_file)
                    self.assertIn(captured.exception.__cause__.winerror, {5, 32, 33})
                    self.assertEqual(locked_file.read(), b"existing output")
                self.assertEqual(len(list(output_file.parent.glob(".output-*.tmp"))), 1)
            finally:
                workbook.close()

    def test_styles_are_registered_in_output_workbook(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            source_file = folder / "source.xlsx"
            output_file = folder / "output.xlsx"

            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "Data"
            cell = worksheet["A1"]
            cell.value = 12.5
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(fill_type="solid", fgColor="4F81BD")
            cell.border = Border(bottom=Side(style="thin", color="000000"))
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )
            cell.number_format = "0.00"
            cell.protection = Protection(locked=False)
            workbook.save(source_file)
            workbook.close()

            merge_workbooks([source_file], output_file)

            merged_workbook = load_workbook(output_file)
            try:
                merged_cell = merged_workbook["Data"]["A1"]
                self.assertEqual(merged_cell.value, 12.5)
                self.assertTrue(merged_cell.font.bold)
                self.assertEqual(merged_cell.fill.fill_type, "solid")
                self.assertEqual(merged_cell.border.bottom.style, "thin")
                self.assertEqual(merged_cell.alignment.horizontal, "center")
                self.assertTrue(merged_cell.alignment.wrap_text)
                self.assertEqual(merged_cell.number_format, "0.00")
                self.assertFalse(merged_cell.protection.locked)
            finally:
                merged_workbook.close()


if __name__ == "__main__":
    unittest.main()
