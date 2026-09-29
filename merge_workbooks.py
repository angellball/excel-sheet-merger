#!/usr/bin/env python3
"""Merge same-named worksheets from Excel and CSV files in a folder."""

from __future__ import annotations

import argparse
import csv
import re
import shutil
import sys
import tempfile
import time
from collections import deque
from concurrent.futures import Future, ThreadPoolExecutor
from copy import copy
from datetime import datetime, timezone
from pathlib import Path
from typing import Callable, Iterable, Iterator, Sequence, TextIO
from zipfile import ZIP_DEFLATED, ZipFile

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.cell.cell import Cell, MergedCell
    from openpyxl.formula.tokenizer import Token, Tokenizer, TokenizerError
    from openpyxl.formula.translate import Translator, TranslatorError
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.worksheet import Worksheet
    from openpyxl.writer.excel import ExcelWriter
except ModuleNotFoundError as error:
    if error.name == "openpyxl":
        raise SystemExit(
            "缺少依赖 openpyxl，请先运行："
            "python3 -m pip install -r requirements.txt"
        ) from None
    raise

try:
    import xlrd
except ModuleNotFoundError as error:
    if error.name != "xlrd":
        raise
    xlrd = None

SUPPORTED_SUFFIXES = {".xlsx", ".xlsm", ".xls", ".csv"}
DEFAULT_WORKERS = 1
DEFAULT_DATETIME_FORMAT = "yyyy-mm-dd hh:mm:ss"
ProgressCallback = Callable[[float, str], None]


def natural_sort_key(value: str) -> tuple[tuple[int, int | str], ...]:
    """Return a case-insensitive key that compares digit groups numerically."""
    return tuple(
        (0, int(part)) if part.isdigit() else (1, part.casefold())
        for part in re.split(r"(\d+)", value)
    )


class TerminalProgress:
    """Render a rate-limited, dependency-free terminal progress bar."""

    def __init__(
        self,
        enabled: bool = True,
        stream: TextIO | None = None,
    ) -> None:
        self.stream = stream or sys.stdout
        self.enabled = enabled and self.stream.isatty()
        self._last_update = 0.0
        self._last_line_length = 0

    def update(self, fraction: float, label: str) -> None:
        if not self.enabled:
            return

        fraction = min(1.0, max(0.0, fraction))
        current_time = time.monotonic()
        if fraction < 1.0 and current_time - self._last_update < 0.08:
            return
        self._last_update = current_time

        terminal_width = shutil.get_terminal_size((100, 20)).columns
        bar_width = max(10, min(36, terminal_width - 45))
        filled_width = round(bar_width * fraction)
        bar = "#" * filled_width + "-" * (bar_width - filled_width)
        label_width = max(0, terminal_width - bar_width - 13)
        if len(label) > label_width:
            label = label[: max(0, label_width - 3)] + "..."
        line = f"[{bar}] {fraction * 100:6.2f}% {label}"
        padding = " " * max(0, self._last_line_length - len(line))
        self.stream.write(f"\r{line}{padding}")
        self.stream.flush()
        self._last_line_length = len(line)

    def close(self) -> None:
        if self.enabled:
            self.stream.write("\n")
            self.stream.flush()


def find_workbooks(folder: Path, output_file: Path) -> list[Path]:
    """Find Excel and CSV files, excluding temporary and output files."""
    output_file = output_file.resolve()
    return sorted(
        (
            path
            for path in folder.rglob("*")
            if path.is_file()
            and path.suffix.lower() in SUPPORTED_SUFFIXES
            and not path.name.startswith("~$")
            and not path.name.startswith(f".{output_file.stem}-")
            and path.resolve() != output_file
        ),
        key=lambda path: natural_sort_key(path.relative_to(folder).as_posix()),
    )


def _safe_sheet_title(title: str) -> str:
    """Return a valid, deterministic Excel worksheet title."""
    title = re.sub(r"[\\/*?:\[\]]", "_", title).strip() or "Sheet1"
    return title[:31]


def _csv_workbook(csv_file: Path) -> Workbook:
    """Load a CSV file into a workbook containing one worksheet."""
    for encoding in ("utf-8-sig", "gb18030"):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = _safe_sheet_title(csv_file.stem)
        try:
            with open(csv_file, "r", encoding=encoding, newline="") as csv_handle:
                sample = csv_handle.read(8192).replace("\r\n", "\n").replace("\r", "\n")
                csv_handle.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
                except csv.Error:
                    dialect = csv.excel
                for row_index, row in enumerate(
                    csv.reader(csv_handle, dialect),
                    start=1,
                ):
                    for column_index, value in enumerate(row, start=1):
                        if value:
                            target_cell = worksheet.cell(row_index, column_index, value)
                            target_cell.data_type = "s"
            return workbook
        except UnicodeDecodeError:
            workbook.close()
        except Exception:
            workbook.close()
            raise
    raise RuntimeError(f"无法按 UTF-8 或 GB18030 解码 CSV 文件：{csv_file}")


def _xls_cell_value(book: xlrd.book.Book, cell: xlrd.sheet.Cell) -> object:
    """Convert an xlrd cell to a value accepted by openpyxl."""
    cell_type = cell.ctype
    if cell_type in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
        return None
    if cell_type == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    if cell_type == xlrd.XL_CELL_DATE:
        date_value = xlrd.xldate_as_datetime(cell.value, book.datemode)
        return date_value.time() if 0 <= cell.value < 1 else date_value
    if cell_type == xlrd.XL_CELL_NUMBER:
        numeric_value = float(cell.value)
        return int(numeric_value) if numeric_value.is_integer() else numeric_value
    if cell_type == xlrd.XL_CELL_ERROR:
        return xlrd.error_text_from_code.get(int(cell.value), "#VALUE!")
    return cell.value


def _xls_workbook(xls_file: Path) -> Workbook:
    """Load a legacy .xls file through xlrd into an openpyxl workbook."""
    if xlrd is None:
        raise RuntimeError(
            "读取 .xls 文件需要依赖 xlrd，请先运行："
            "python -m pip install -r requirements.txt"
        )

    source_workbook = xlrd.open_workbook(
        str(xls_file),
        on_demand=True,
        formatting_info=True,
    )
    workbook = Workbook()
    workbook.remove(workbook.active)
    try:
        for source_sheet in source_workbook.sheets():
            worksheet = workbook.create_sheet(source_sheet.name)
            for row_index in range(source_sheet.nrows):
                for column_index, cell in enumerate(
                    source_sheet.row(row_index),
                    start=1,
                ):
                    value = _xls_cell_value(source_workbook, cell)
                    if value is not None:
                        target_cell = worksheet.cell(
                            row=row_index + 1,
                            column=column_index,
                            value=value,
                        )
                        if cell.ctype == xlrd.XL_CELL_TEXT:
                            target_cell.data_type = "s"
                        source_format = source_workbook.xf_list[cell.xf_index]
                        number_format = source_workbook.format_map.get(
                            source_format.format_key,
                        )
                        if (
                            number_format is not None
                            and isinstance(number_format.format_str, str)
                            and number_format.format_str.strip()
                        ):
                            target_cell.number_format = number_format.format_str
                        elif target_cell.data_type == "d":
                            target_cell.number_format = DEFAULT_DATETIME_FORMAT
            for merged_range in source_sheet.merged_cells:
                first_row, last_row, first_column, last_column = merged_range
                worksheet.merge_cells(
                    start_row=first_row + 1,
                    end_row=last_row,
                    start_column=first_column + 1,
                    end_column=last_column,
                )
            for column_index, column_info in source_sheet.colinfo_map.items():
                worksheet.column_dimensions[get_column_letter(column_index + 1)].width = (
                    column_info.width / 256
                )
            for row_index, row_info in source_sheet.rowinfo_map.items():
                if not row_info.has_default_height:
                    worksheet.row_dimensions[row_index + 1].height = row_info.height / 20
    except Exception:
        workbook.close()
        raise
    finally:
        source_workbook.release_resources()
    return workbook


def worksheet_has_content(worksheet: Worksheet) -> bool:
    """Check whether a worksheet contains at least one non-empty cell."""
    return any(
        not isinstance(cell, MergedCell) and cell.value is not None
        for cell in worksheet._cells.values()
    )


CELL_REFERENCE_RE = re.compile(
    r"(?P<column>\$?[A-Za-z]{1,3})(?P<row>\$?[1-9][0-9]{0,6})$"
)
ROW_REFERENCE_RE = re.compile(
    r"(?P<start>\$?[1-9][0-9]{0,6}):(?P<end>\$?[1-9][0-9]{0,6})$"
)


def _formula_sheet_name(sheet_prefix: str) -> str | None:
    if not sheet_prefix:
        return None

    sheet_name = sheet_prefix[:-1]
    if "]" in sheet_name:
        sheet_name = sheet_name.rsplit("]", 1)[1]
    if sheet_name.startswith("'") and sheet_name.endswith("'"):
        sheet_name = sheet_name[1:-1].replace("''", "'")
    return sheet_name


def _formula_reference_matches_sheet(
    sheet_prefix: str,
    source_sheet_name: str,
) -> bool:
    referenced_sheet_name = _formula_sheet_name(sheet_prefix)
    return referenced_sheet_name is None or (
        referenced_sheet_name.casefold() == source_sheet_name.casefold()
    )


def _translate_formula_reference_part(
    reference_part: str,
    row_mapping: dict[int, int],
) -> str:
    cell_match = CELL_REFERENCE_RE.fullmatch(reference_part)
    if cell_match:
        source_row_text = cell_match.group("row")
        if source_row_text.startswith("$"):
            return reference_part
        source_row = int(source_row_text)
        target_row = row_mapping.get(source_row)
        if target_row is None:
            return reference_part
        return f"{cell_match.group('column')}{target_row}"

    row_match = ROW_REFERENCE_RE.fullmatch(reference_part)
    if row_match:
        translated_rows: list[str] = []
        for row_text in (row_match.group("start"), row_match.group("end")):
            if row_text.startswith("$"):
                translated_rows.append(row_text)
                continue
            target_row = row_mapping.get(int(row_text))
            translated_rows.append(str(target_row) if target_row is not None else row_text)
        return ":".join(translated_rows)

    return reference_part


def translate_formula_rows(
    formula: str,
    row_mapping: dict[int, int],
    source_sheet_name: str,
) -> str:
    """Translate relative row references using the rows retained in output."""
    translated_formula: list[str] = []
    for token in Tokenizer(formula).items:
        if token.type != Token.OPERAND or token.subtype != Token.RANGE:
            translated_formula.append(token.value)
            continue

        sheet_prefix = ""
        reference_text = token.value
        if "!" in reference_text:
            sheet_prefix, reference_text = reference_text.rsplit("!", 1)
            sheet_prefix += "!"
        if not _formula_reference_matches_sheet(sheet_prefix, source_sheet_name):
            translated_formula.append(token.value)
            continue

        translated_parts = [
            _translate_formula_reference_part(reference_part, row_mapping)
            for reference_part in reference_text.split(":")
        ]
        translated_formula.append(sheet_prefix + ":".join(translated_parts))

    return "=" + "".join(translated_formula)


def source_row_has_content(
    source_row: Sequence[Cell | MergedCell],
    merged_rows: set[int],
) -> bool:
    """Return whether a source row should occupy a row in the output."""
    return source_row[0].row in merged_rows or any(
        not isinstance(source_cell, MergedCell)
        and (
            source_cell.value is not None
            or source_cell.hyperlink is not None
            or source_cell.comment is not None
        )
        for source_cell in source_row
    )


def copy_cell(
    source: Cell,
    target: Cell,
    style_cache: dict[int, object],
    row_mapping: dict[int, int] | None = None,
    source_sheet_name: str | None = None,
) -> None:
    """Copy a cell value and its common presentation properties."""
    if source.data_type == "f" and isinstance(source.value, str):
        if row_mapping is not None and source_sheet_name is not None:
            try:
                translated_formula = translate_formula_rows(
                    source.value,
                    row_mapping,
                    source_sheet_name,
                )
            except TokenizerError:
                translated_formula = source.value
        else:
            try:
                translated_formula = Translator(
                    source.value,
                    origin=source.coordinate,
                ).translate_formula(target.coordinate)
            except (TokenizerError, TranslatorError):
                translated_formula = source.value

        target.value = (
            source.value
            if "#REF!" in translated_formula and "#REF!" not in source.value
            else translated_formula
        )
    else:
        target.value = source.value
        target.data_type = source.data_type
    if source.has_style:
        source_style_id = source.style_id
        number_format = source.number_format
        valid_number_format = (
            isinstance(number_format, str) and bool(number_format.strip())
        )
        cached_style = style_cache.get(source_style_id) if valid_number_format else None
        if cached_style is None:
            # Style IDs are workbook-local. Register each distinct style once,
            # then reuse the translated destination style for matching cells.
            target.font = copy(source.font)
            target.fill = copy(source.fill)
            target.border = copy(source.border)
            target.alignment = copy(source.alignment)
            if valid_number_format:
                target.number_format = number_format
            elif target.data_type == "d":
                target.number_format = DEFAULT_DATETIME_FORMAT
            target.protection = copy(source.protection)
            target.quotePrefix = source.quotePrefix
            target.pivotButton = source.pivotButton
            cached_style = copy(target._style)
            if valid_number_format:
                style_cache[source_style_id] = cached_style
        else:
            target._style = copy(cached_style)
    if source.hyperlink:
        target._hyperlink = copy(source.hyperlink)
    if source.comment:
        target.comment = copy(source.comment)


def copy_column_widths(source: Worksheet, target: Worksheet) -> None:
    """Keep the greatest width found for each target column."""
    for column_index in range(1, source.max_column + 1):
        column_letter = get_column_letter(column_index)
        source_dimension = source.column_dimensions.get(column_letter)
        if source_dimension is None or source_dimension.width is None:
            continue

        source_width = source_dimension.width
        target_dimension = target.column_dimensions.get(column_letter)
        target_width = target_dimension.width if target_dimension else None
        if source_width is not None and (
            target_width is None or source_width > target_width
        ):
            target.column_dimensions[column_letter].width = source_width


def copy_worksheet_rows(
    source: Worksheet,
    target: Worksheet,
    target_start_row: int,
    source_start_row: int,
    style_cache: dict[int, object],
    row_progress: Callable[[int, int], None] | None = None,
) -> int:
    """Append source rows to target and return the next available target row."""
    total_rows = source.max_row - source_start_row + 1
    source_rows = source.iter_rows(
        min_row=source_start_row,
        max_row=source.max_row,
        max_col=source.max_column,
    )

    merged_rows: set[int] = set()
    for merged_range in source.merged_cells.ranges:
        if merged_range.min_row < source_start_row:
            continue
        merged_cell = source.cell(
            row=merged_range.min_row,
            column=merged_range.min_col,
        )
        if merged_cell.value is not None:
            merged_rows.update(range(merged_range.min_row, merged_range.max_row + 1))

    kept_source_rows = [
        source_row[0].row
        for source_row in source_rows
        if source_row_has_content(source_row, merged_rows)
    ]
    source_to_target_row = {
        source_row_number: target_start_row + row_index
        for row_index, source_row_number in enumerate(kept_source_rows)
    }
    for source_row_number in range(1, source_start_row):
        source_to_target_row.setdefault(source_row_number, source_row_number)

    target_row_number = target_start_row
    source_rows = source.iter_rows(
        min_row=source_start_row,
        max_row=source.max_row,
        max_col=source.max_column,
    )
    for row_index, source_row in enumerate(source_rows, start=1):
        if not source_row_has_content(source_row, merged_rows):
            if row_progress and (
                row_index == 1 or row_index % 100 == 0 or row_index == total_rows
            ):
                row_progress(row_index, total_rows)
            continue

        source_row_number = source_row[0].row
        source_to_target_row[source_row_number] = target_row_number
        for source_cell in source_row:
            if isinstance(source_cell, MergedCell):
                continue
            if (
                source_cell.value is None
                and not source_cell.has_style
                and source_cell.hyperlink is None
                and source_cell.comment is None
            ):
                continue
            target_cell = target.cell(
                row=target_row_number,
                column=source_cell.column,
            )
            copy_cell(
                source_cell,
                target_cell,
                style_cache,
                row_mapping=source_to_target_row,
                source_sheet_name=source.title,
            )

        source_height = source.row_dimensions[source_row_number].height
        if source_height is not None:
            target.row_dimensions[target_row_number].height = source_height

        if row_progress and (
            row_index == 1 or row_index % 100 == 0 or row_index == total_rows
        ):
            row_progress(row_index, total_rows)

        target_row_number += 1

    for merged_range in source.merged_cells.ranges:
        if (
            merged_range.min_row < source_start_row
            or merged_range.min_row not in source_to_target_row
            or merged_range.max_row not in source_to_target_row
        ):
            continue
        target.merge_cells(
            start_row=source_to_target_row[merged_range.min_row],
            start_column=merged_range.min_col,
            end_row=source_to_target_row[merged_range.max_row],
            end_column=merged_range.max_col,
        )

    copy_column_widths(source, target)
    return target_row_number


def load_workbook_file(workbook_file: Path) -> Workbook:
    """Load one supported input file and add its path to any read error."""
    try:
        suffix = workbook_file.suffix.casefold()
        if suffix == ".csv":
            return _csv_workbook(workbook_file)
        if suffix == ".xls":
            return _xls_workbook(workbook_file)
        return load_workbook(workbook_file, data_only=False)
    except Exception as error:
        raise RuntimeError(f"无法读取文件：{workbook_file}\n原因：{error}") from error


def iter_loaded_workbooks(
    workbook_files: Sequence[Path],
    workers: int,
) -> Iterator[tuple[Path, Workbook]]:
    """Prefetch workbooks concurrently while yielding them in source order."""
    if workers < 1:
        raise ValueError("workers must be at least 1")
    if workers == 1 or len(workbook_files) <= 1:
        for workbook_file in workbook_files:
            yield workbook_file, load_workbook_file(workbook_file)
        return

    pending: deque[tuple[Path, Future[Workbook]]] = deque()
    file_iterator = iter(workbook_files)
    with ThreadPoolExecutor(
        max_workers=min(workers, len(workbook_files)),
        thread_name_prefix="excel-loader",
    ) as executor:
        for _ in range(min(workers, len(workbook_files))):
            workbook_file = next(file_iterator)
            pending.append(
                (workbook_file, executor.submit(load_workbook_file, workbook_file))
            )

        try:
            while pending:
                workbook_file, future = pending.popleft()
                workbook = future.result()
                try:
                    next_file = next(file_iterator)
                except StopIteration:
                    pass
                else:
                    pending.append(
                        (next_file, executor.submit(load_workbook_file, next_file))
                    )
                yield workbook_file, workbook
        finally:
            for _, future in pending:
                if future.cancel():
                    continue
                try:
                    future.result().close()
                except Exception:
                    pass


def save_workbook_atomic(workbook: Workbook, output_file: Path) -> None:
    """Write to a temporary file, replacing the output only after success."""
    output_file.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{output_file.stem}-",
            suffix=".tmp",
            dir=output_file.parent,
            delete=False,
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)
            with ZipFile(
                temporary_file,
                mode="w",
                compression=ZIP_DEFLATED,
                allowZip64=True,
            ) as archive:
                workbook.properties.modified = datetime.now(timezone.utc).replace(
                    tzinfo=None,
                )
                ExcelWriter(workbook, archive).write_data()
    except Exception:
        if temporary_path is not None:
            try:
                temporary_path.unlink(missing_ok=True)
            except OSError:
                pass
        raise

    for attempt in range(6):
        try:
            temporary_path.replace(output_file)
            return
        except OSError as error:
            if getattr(error, "winerror", None) in {5, 32, 33} and attempt < 5:
                time.sleep(0.1 * 2**attempt)
                continue
            raise RuntimeError(
                f"无法替换结果文件：{output_file}\n原因：{error}\n"
                "请关闭 Excel/WPS 等程序中打开的结果文件，或改用其他输出路径。\n"
                f"本次完整合并结果已保留在：{temporary_path}\n"
                "解除占用后，可将该临时文件改名为 .xlsx 文件。"
            ) from error


def merge_workbooks(
    workbook_files: Iterable[Path],
    output_file: Path,
    keep_all_headers: bool = False,
    header_rows: int = 1,
    workers: int = DEFAULT_WORKERS,
    progress: ProgressCallback | None = None,
) -> tuple[int, int]:
    """Merge workbooks and return the workbook and worksheet copy counts."""
    if header_rows < 0:
        raise ValueError("header_rows cannot be negative")
    if workers < 1:
        raise ValueError("workers must be at least 1")

    workbook_files = list(workbook_files)
    total_files = len(workbook_files)
    output_workbook = Workbook()
    output_workbook.remove(output_workbook.active)
    next_rows: dict[str, int] = {}
    workbook_count = 0
    worksheet_count = 0

    if progress and workbook_files:
        progress(0.0, f"正在读取 {workbook_files[0].name}")

    try:
        loaded_workbooks = iter_loaded_workbooks(workbook_files, workers)
        for file_index, (workbook_file, workbook) in enumerate(
            loaded_workbooks,
            start=1,
        ):
            workbook_count += 1
            style_cache: dict[int, object] = {}
            try:
                source_sheets = workbook.worksheets
                sheet_count = max(1, len(source_sheets))
                for sheet_index, source_sheet in enumerate(source_sheets, start=1):
                    if not worksheet_has_content(source_sheet):
                        if progress and total_files:
                            completed = (
                                file_index - 1 + sheet_index / sheet_count
                            ) / total_files
                            progress(
                                completed * 0.9,
                                f"{workbook_file.name} / {source_sheet.title}（空表）",
                            )
                        continue

                    sheet_name = source_sheet.title
                    if sheet_name not in output_workbook.sheetnames:
                        target_sheet = output_workbook.create_sheet(sheet_name)
                        next_rows[sheet_name] = 1
                        if source_sheet.freeze_panes:
                            freeze_panes = source_sheet.freeze_panes
                            target_sheet.freeze_panes = getattr(
                                freeze_panes,
                                "coordinate",
                                freeze_panes,
                            )
                    else:
                        target_sheet = output_workbook[sheet_name]

                    source_start_row = 1
                    if next_rows[sheet_name] > 1 and not keep_all_headers:
                        source_start_row = header_rows + 1

                    if source_start_row <= source_sheet.max_row:

                        def report_rows(current_row: int, total_rows: int) -> None:
                            if not progress or not total_files:
                                return
                            sheet_fraction = current_row / max(1, total_rows)
                            file_fraction = (
                                sheet_index - 1 + sheet_fraction
                            ) / sheet_count
                            completed = (
                                file_index - 1 + file_fraction
                            ) / total_files
                            progress(
                                completed * 0.9,
                                f"{workbook_file.name} / {source_sheet.title} "
                                f"{current_row}/{total_rows}",
                            )

                        next_rows[sheet_name] = copy_worksheet_rows(
                            source_sheet,
                            target_sheet,
                            next_rows[sheet_name],
                            source_start_row,
                            style_cache,
                            row_progress=report_rows,
                        )
                        worksheet_count += 1
            finally:
                workbook.close()

            if progress and total_files:
                progress(
                    file_index / total_files * 0.9,
                    f"已处理 {file_index}/{total_files}：{workbook_file.name}",
                )

        if not output_workbook.sheetnames:
            output_workbook.create_sheet("合并结果")

        if progress:
            progress(0.92, "正在写入结果文件")
        save_workbook_atomic(output_workbook, output_file)
        if progress:
            progress(1.0, "完成")
    finally:
        output_workbook.close()

    return workbook_count, worksheet_count


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="把文件夹内 Excel/CSV 文件中的同名工作表纵向合并到一个文件中。"
    )
    parser.add_argument(
        "folder",
        nargs="?",
        type=Path,
        help="待合并 Excel/CSV 文件所在文件夹（省略时运行后输入）",
    )
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        help="输出文件路径（默认：输入文件夹/合并结果.xlsx）",
    )
    parser.add_argument(
        "--header-rows",
        type=int,
        help="标题行数；输入 0 表示没有标题（交互运行时会提示，默认：1）",
    )
    parser.add_argument(
        "--keep-all-headers",
        action="store_true",
        help="保留每个来源工作表的标题，不跳过重复标题",
    )
    parser.add_argument(
        "--workers",
        type=int,
        default=DEFAULT_WORKERS,
        help=f"并发预读文件数（默认：{DEFAULT_WORKERS}；内存不足可设为 1）",
    )
    parser.add_argument(
        "--no-progress",
        action="store_true",
        help="不显示进度条",
    )
    args = parser.parse_args()
    if args.header_rows is not None and args.header_rows < 0:
        parser.error("--header-rows 不能小于 0")
    if args.workers < 1:
        parser.error("--workers 不能小于 1")
    return args


def get_input_folder(folder_argument: Path | None) -> Path:
    """Resolve the CLI folder argument or prompt for one interactively."""
    if folder_argument is None:
        try:
            entered_folder = input(
                "请输入要扫描的 Excel/CSV 文件夹路径（留空使用当前文件夹）："
            ).strip()
        except EOFError:
            raise SystemExit("错误：未输入文件夹路径") from None

        if (
            len(entered_folder) >= 2
            and entered_folder[0] == entered_folder[-1]
            and entered_folder[0] in {"'", '"'}
        ):
            entered_folder = entered_folder[1:-1]
        folder_argument = Path(entered_folder) if entered_folder else Path.cwd()

    return folder_argument.expanduser().resolve()


def get_header_rows(
    header_rows_argument: int | None,
    prompt_when_missing: bool,
) -> int:
    """Resolve the title row count, prompting during interactive use."""
    if header_rows_argument is not None:
        return header_rows_argument
    if not prompt_when_missing:
        return 1

    while True:
        try:
            entered_rows = input(
                "请输入标题行数（默认 1，输入 0 表示没有标题行）："
            ).strip()
        except EOFError:
            raise SystemExit("错误：未输入标题行数") from None

        if not entered_rows:
            return 1
        try:
            header_rows = int(entered_rows)
        except ValueError:
            print("标题行数必须是 0 或正整数，请重新输入。")
            continue
        if header_rows >= 0:
            return header_rows
        print("标题行数不能小于 0，请重新输入。")


def main() -> None:
    args = parse_args()
    interactive = args.folder is None
    folder = get_input_folder(args.folder)
    if not folder.is_dir():
        raise SystemExit(f"错误：文件夹不存在：{folder}")

    header_rows = get_header_rows(
        args.header_rows,
        prompt_when_missing=interactive and not args.keep_all_headers,
    )

    output_file = (
        args.output.expanduser().resolve()
        if args.output
        else folder / "合并结果.xlsx"
    )
    workbook_files = find_workbooks(folder, output_file)
    if not workbook_files:
        raise SystemExit(
            f"错误：{folder} 及其子文件夹中没有可合并的 .xls、.xlsx、.xlsm 或 .csv 文件"
        )

    if args.keep_all_headers:
        header_message = "保留每个文件的标题"
    else:
        header_message = f"后续同名工作表跳过 {header_rows} 行标题"
    print(f"找到 {len(workbook_files)} 个 Excel/CSV 文件，已按名称自然升序排列。")
    print(f"{header_message}，预读并发数 {args.workers}，开始合并……")

    progress_bar = TerminalProgress(enabled=not args.no_progress)
    try:
        workbook_count, worksheet_count = merge_workbooks(
            workbook_files,
            output_file,
            keep_all_headers=args.keep_all_headers,
            header_rows=header_rows,
            workers=args.workers,
            progress=progress_bar.update,
        )
    except (OSError, RuntimeError, ValueError, TypeError) as error:
        progress_bar.close()
        raise SystemExit(f"错误：{error}") from None
    progress_bar.close()
    print(
        f"合并完成：读取 {workbook_count} 个文件，"
        f"合并 {worksheet_count} 个工作表，输出到：{output_file}"
    )


if __name__ == "__main__":
    main()
